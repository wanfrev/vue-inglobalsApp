"""Extracción de texto de archivos que el usuario adjunta a una consulta
puntual en el chat (evidencia de esa pregunta, no se indexan en la base de
leyes compartida — para eso está app.core.rag / POST /api/v1/documents).

Formatos soportados: PDF, TXT, Word (.docx) y Excel (.xlsx). Imágenes o PDFs
escaneados (sin texto extraíble) no se soportan todavía — requieren OCR
(Tesseract), que es una dependencia de sistema, no solo una librería Python."""

import io

import openpyxl
from docx import Document
from PyPDF2 import PdfReader

SUPPORTED_EXTENSIONS = (".pdf", ".txt", ".docx", ".xlsx")
MAX_ATTACHMENTS_PER_REQUEST = 5


def _extract_pdf(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    pages = [page.extract_text() or "" for page in reader.pages]
    text = "\n\n".join(p for p in pages if p.strip())
    if not text.strip():
        raise ValueError(
            "El PDF no tiene texto extraíble (parece ser un escaneo/imagen). "
            "Por ahora solo se soportan PDFs con texto real."
        )
    return text


def _extract_docx(content: bytes) -> str:
    doc = Document(io.BytesIO(content))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def _extract_xlsx(content: bytes) -> str:
    workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    parts = []

    for sheet in workbook.worksheets:
        parts.append(f"[Hoja: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            values = [str(v) for v in row if v is not None]
            if values:
                parts.append(" | ".join(values))

    return "\n".join(parts)


def extract_text(filename: str, content: bytes) -> str:
    lowered = filename.lower()

    try:
        if lowered.endswith(".pdf"):
            return _extract_pdf(content)

        if lowered.endswith(".txt"):
            return content.decode("utf-8", errors="ignore")

        if lowered.endswith(".docx"):
            return _extract_docx(content)

        if lowered.endswith(".xlsx"):
            return _extract_xlsx(content)
    except ValueError:
        raise
    except Exception:
        raise ValueError(
            f"No se pudo leer '{filename}'. ¿Está dañado o el contenido no corresponde a su extensión?"
        )

    raise ValueError(
        f"Formato no soportado para '{filename}'. Se aceptan: PDF, TXT, Word (.docx) o Excel (.xlsx)."
    )
