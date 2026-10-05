import json
import tempfile
import uuid
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
from PyPDF2 import PdfReader
from docx import Document as DocxDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter

from app.config import settings
from app.core import lexical
from app.database import (
    faiss_write_lock,
    get_embedding_model,
    get_faiss_index,
    get_metadata,
    refresh_faiss_from_disk,
    save_faiss_index,
)


def load_pdf_text(file_path: Path) -> str:
    reader = PdfReader(str(file_path))
    pages = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            pages.append(text)
    return "\n\n".join(pages)


def load_docx_text(file_path: Path) -> str:
    doc = DocxDocument(str(file_path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)
    return "\n\n".join(paragraphs)


def load_document_text(file_path: Path) -> str:
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        return load_pdf_text(file_path)
    if suffix == ".docx":
        return load_docx_text(file_path)
    if suffix == ".txt":
        return file_path.read_text(encoding="utf-8")
    raise ValueError(f"Tipo de archivo no soportado para indexar: '{file_path.suffix}'")


def extract_attachment_text(filename: str, content: bytes, max_chars: int | None = None) -> str:
    """Extrae el texto de un archivo adjunto a UNA consulta puntual del chat
    (a pedido del cliente) — a diferencia de index_document(), esto NO se
    guarda ni se indexa en la bibliografía compartida: es contexto efímero
    que solo se usa para responder esa consulta (ver `_run_loop1` en
    engine.py). Mismos formatos que la bibliografía (PDF/Word/texto plano),
    reutilizando load_document_text(). Si el archivo es una imagen escaneada
    sin texto real, no hay OCR todavía — se informa al usuario en vez de
    fallar en silencio con un contexto vacío."""
    suffix = Path(filename).suffix.lower()
    if suffix not in (".pdf", ".docx", ".txt"):
        raise ValueError(
            f"Formato de archivo no soportado: '{suffix or 'desconocido'}'. "
            "Se aceptan PDF, Word (.docx) o texto plano (.txt)."
        )

    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
        tmp.write(content)
        tmp_path = Path(tmp.name)

    try:
        text = load_document_text(tmp_path).strip()
    finally:
        tmp_path.unlink(missing_ok=True)

    if not text:
        raise ValueError(
            "No se pudo extraer texto del archivo — ¿es una imagen escaneada sin texto "
            "real? Por ahora no se soporta reconocimiento óptico (OCR)."
        )

    limit = max_chars or settings.ATTACHMENT_MAX_CHARS
    if len(text) > limit:
        text = text[:limit] + "\n[...documento truncado por longitud...]"

    return text


def chunk_text(
    text: str, chunk_size: int = 1000, overlap: int = 200
) -> list[str]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=overlap,
        length_function=len,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    return [chunk.strip() for chunk in splitter.split_text(text) if chunk.strip()]


def index_document(
    file_path: Path,
    title: str,
    category: str,
    date: str | None = None,
) -> dict:
    text = load_document_text(file_path)
    if not text.strip():
        raise ValueError(f"El PDF '{file_path.name}' no contiene texto extraíble")

    chunks = chunk_text(text)
    if not chunks:
        raise ValueError(f"No se pudieron generar fragmentos desde '{file_path.name}'")

    model = get_embedding_model()
    embeddings = model.encode(
        chunks,
        convert_to_numpy=True,
        normalize_embeddings=True,
        show_progress_bar=True,
    )

    doc_id = str(uuid.uuid4())
    indexed_at = datetime.now(timezone.utc).isoformat()

    # La codificación (lenta) ya terminó fuera del candado; solo la parte de
    # leer-modificar-guardar es exclusiva. Se recarga antes desde disco para
    # no pisar lo que otro worker haya subido mientras tanto.
    with faiss_write_lock():
        refresh_faiss_from_disk(dimension=embeddings.shape[1])
        index = get_faiss_index(dimension=embeddings.shape[1])
        metadata = get_metadata()

        for i, (chunk, vector) in enumerate(zip(chunks, embeddings)):
            metadata.append(
                {
                    "id": doc_id,
                    "chunk_index": i,
                    "title": title,
                    "category": category,
                    "date": date or "",
                    "text": chunk,
                    "indexed_at": indexed_at,
                    "file_name": file_path.name,
                }
            )

        index.add(np.array(embeddings, dtype=np.float32))
        save_faiss_index()

    return {
        "id": doc_id,
        "title": title,
        "category": category,
        "chunks": len(chunks),
        "indexed_at": indexed_at,
    }


def _expand_full_documents(results: list[dict], metadata: list[dict], max_doc_chars: int = 25000) -> list[dict]:
    """Un artículo puede remitir a otros del MISMO documento (ej. el art. 23 de
    la Providencia SNAT/2024/000102 remite a los art. 3, 4, 5 y 19) — si la
    búsqueda semántica por top_k trae solo el chunk que menciona la remisión,
    el Loop 1 ve la referencia pero no el texto de los artículos referidos, y
    reporta (con razón) que "falta el texto completo" aunque esa bibliografía
    SÍ está indexada. Por eso, una vez que un documento entra al top_k, se
    completa con el resto de sus chunks (en orden), no solo el fragmento que
    matcheó semánticamente. Acotado por tamaño total del documento
    (max_doc_chars, ~una Providencia o norma corta/mediana completa) para no
    inflar el contexto — y el costo — con libros o manuales largos que solo
    aportaron un fragmento puntual."""
    seen_keys = {(r["id"], r["chunk_index"]) for r in results}
    seen_doc_ids = {r["id"] for r in results}
    chunks_by_doc: dict[str, list[dict]] = {}
    for m in metadata:
        doc_id = m.get("id")
        if doc_id in seen_doc_ids:
            chunks_by_doc.setdefault(doc_id, []).append(m)

    expansions = []
    for doc_id, chunks in chunks_by_doc.items():
        total_chars = sum(len(c.get("text", "")) for c in chunks)
        if total_chars > max_doc_chars:
            continue
        for m in sorted(chunks, key=lambda c: c.get("chunk_index", 0)):
            key = (m["id"], m["chunk_index"])
            if key in seen_keys:
                continue
            seen_keys.add(key)
            entry = m.copy()
            entry["score"] = 0.0
            expansions.append(entry)

    return results + expansions


# Candidatos que aporta cada búsqueda (semántica y léxica) antes de fusionarlas.
CANDIDATE_POOL = 40
# Constante de la fusión por rango recíproco (RRF); 60 es el valor estándar.
RRF_K = 60


def search_legal_context(
    query: str,
    top_k: int = 5,
    filter_category: str | None = None,
    exclude_categories: list[str] | None = None,
    expand_full_document: bool = True,
    hybrid: bool = True,
) -> list[dict]:
    """Búsqueda híbrida: semántica (FAISS + MiniLM) fusionada con léxica (BM25,
    ver lexical.py) por rango recíproco. Solo la semántica fallaba con una
    bibliografía grande en español: consultas con términos legales exactos
    ("retención de IVA", "contribuyentes especiales") no traían la
    providencia que los regula aunque estuviera indexada."""
    model = get_embedding_model()
    query_embedding = model.encode(
        [query], convert_to_numpy=True, normalize_embeddings=True
    )

    index = get_faiss_index()
    metadata = get_metadata()

    if index.ntotal == 0:
        return []

    allowed = None
    if filter_category or exclude_categories:
        excluded = set(exclude_categories or [])
        allowed = [
            i
            for i, m in enumerate(metadata)
            if (not filter_category or m.get("category") == filter_category)
            and m.get("category") not in excluded
        ]
        if not allowed:
            return []

    pool = max(top_k * 6, CANDIDATE_POOL)

    # --- candidatos semánticos: (posición global, coseno)
    if allowed is None:
        k = min(pool, index.ntotal)
        scores, positions = index.search(query_embedding, k)
        dense = [(int(p), float(sc)) for sc, p in zip(scores[0], positions[0]) if p != -1]
    else:
        import faiss

        allowed_vectors = np.array([index.reconstruct(i) for i in allowed], dtype=np.float32)
        sub_index = faiss.IndexFlatIP(query_embedding.shape[1])
        sub_index.add(allowed_vectors)
        scores, sub_positions = sub_index.search(query_embedding, min(pool, len(allowed)))
        dense = [
            (allowed[int(p)], float(sc))
            for sc, p in zip(scores[0], sub_positions[0])
            if p != -1
        ]
    cosine = dict(dense)

    # --- candidatos léxicos
    lexical_rank: list[int] = []
    if hybrid:
        bm25 = lexical.bm25_scores(metadata, query)
        if allowed is not None:
            masked = np.zeros_like(bm25)
            masked[allowed] = bm25[allowed]
            bm25 = masked
        top = np.argsort(-bm25)[:pool]
        lexical_rank = [int(i) for i in top if bm25[i] > 0]

    # --- fusión por rango recíproco
    fused: dict[int, float] = {}
    for rank_list in ([i for i, _ in dense], lexical_rank):
        for rank, i in enumerate(rank_list):
            fused[i] = fused.get(i, 0.0) + 1.0 / (RRF_K + rank + 1)
    ordered = sorted(fused, key=fused.get, reverse=True)[:top_k]

    results = []
    for i in ordered:
        entry = metadata[i].copy()
        entry["score"] = cosine.get(i, float(index.reconstruct(i) @ query_embedding[0]))
        results.append(entry)

    if expand_full_document and results:
        results = _expand_full_documents(results, metadata)
    return results


def list_documents(filter_category: str | None = None) -> list[dict]:
    metadata = get_metadata()
    docs: dict[str, dict] = {}

    for entry in metadata:
        if filter_category and entry.get("category") != filter_category:
            continue
        doc_id = entry.get("id")
        if doc_id not in docs:
            docs[doc_id] = {
                "id": doc_id,
                "title": entry.get("title", ""),
                "category": entry.get("category", ""),
                "chunks": 0,
                "indexed_at": entry.get("indexed_at", ""),
                "date": entry.get("date", ""),
                "file_name": entry.get("file_name", ""),
            }
        docs[doc_id]["chunks"] += 1

    return sorted(docs.values(), key=lambda d: d["indexed_at"], reverse=True)


def delete_document(doc_id: str) -> None:
    with faiss_write_lock():
        refresh_faiss_from_disk()
        metadata = get_metadata()
        indices_to_remove = [i for i, m in enumerate(metadata) if m.get("id") == doc_id]

        if not indices_to_remove:
            raise ValueError(f"Documento '{doc_id}' no encontrado en el índice")

        index = get_faiss_index()
        all_vectors = np.array(
            [index.reconstruct(i) for i in range(index.ntotal)], dtype=np.float32
        )

        keep_mask = np.ones(index.ntotal, dtype=bool)
        keep_mask[indices_to_remove] = False
        remaining_vectors = all_vectors[keep_mask]

        import faiss

        dimension = remaining_vectors.shape[1] if remaining_vectors.size > 0 else 384
        new_index = faiss.IndexFlatIP(dimension)
        if remaining_vectors.size > 0:
            new_index.add(remaining_vectors)

        import app.database as db

        db._faiss_index = new_index

        removed = set(indices_to_remove)
        remaining_metadata = [m for i, m in enumerate(metadata) if i not in removed]
        metadata.clear()
        metadata.extend(remaining_metadata)

        save_faiss_index()
